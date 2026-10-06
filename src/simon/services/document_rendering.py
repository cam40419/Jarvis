"""Deterministic, semantic Word reports from saved Markdown; no network or model calls."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any, cast
from urllib.parse import urlsplit
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from markdown_it import MarkdownIt
from markdown_it.token import Token

from simon.services.report_layout import prepare_report

RENDERER_VERSION = "1"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_MARKDOWN_CHARS = 1_000_000
_INK = "242424"
_MUTED = "626262"
_ACCENT = "315D8A"
_LINE = "D9D9D5"
_PAPER = "F4F4F1"
_URL = re.compile(r"https?://[^\s<>]+")


def _element(tag: str, **attributes: str) -> Any:
    element = OxmlElement(tag)
    for name, value in attributes.items():
        element.set(qn("w:" + name), value)
    return element


def _rule(paragraph: Paragraph, color: str = _LINE, *, width: str = "6") -> None:
    borders = _element("w:pBdr")
    borders.append(_element("w:bottom", val="single", sz=width, space="7", color=color))
    paragraph._p.get_or_add_pPr().append(borders)


def _field(paragraph: Paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    run._r.append(_element("w:fldChar", fldCharType="begin"))
    text = _element("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = " " + instruction + " "
    run._r.append(text)
    run._r.append(_element("w:fldChar", fldCharType="separate"))
    value = _element("w:t")
    value.text = "1"
    run._r.append(value)
    run._r.append(_element("w:fldChar", fldCharType="end"))


def _font_family(properties: Any, family: str) -> None:
    fonts = properties.get_or_add_rFonts()
    # Word's template theme can override an explicit name in other consumers.
    for attribute in tuple(fonts.attrib):
        if attribute.casefold().endswith("theme"):
            del fonts.attrib[attribute]
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + slot), family)


def _style(document: DocumentType, name: str, size: float) -> ParagraphStyle:
    styles = document.styles
    style = cast(
        ParagraphStyle,
        styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH),
    )
    _font_family(style.element.get_or_add_rPr(), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(_INK)
    style.paragraph_format.widow_control = True
    return style


def _document(title: str, project_name: str) -> DocumentType:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.8), Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.7)
    section.header_distance = section.footer_distance = Inches(0.35)
    normal = _style(document, "Normal", 10.5)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(6)
    heading_sizes = (18, 13.5, 11.5, 11, 10.5, 10.5)
    for level, size in enumerate(heading_sizes, 1):
        style = _style(document, f"Heading {level}", size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if level <= 2 else 9)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    style = _style(document, "Title", 27)
    style.font.bold = True
    style.paragraph_format.space_after = Pt(17)
    style.paragraph_format.keep_with_next = True
    # The title has one explicit rule; avoid the template's inherited theme rule.
    for inherited in style.element.xpath("./w:pPr/w:pBdr | ./w:pPr/w:contextualSpacing"):
        inherited.getparent().remove(inherited)
    code = _style(document, "Report Code", 9)
    _font_family(code.element.get_or_add_rPr(), "Courier New")
    code.paragraph_format.line_spacing = 1.08
    code.paragraph_format.left_indent = Inches(0.13)
    code.paragraph_format.right_indent = Inches(0.13)
    quote = _style(document, "Report Quote", 10.5)
    quote.font.color.rgb = RGBColor.from_string(_MUTED)
    quote.paragraph_format.left_indent = Inches(0.22)
    quote.paragraph_format.right_indent = Inches(0.1)
    quote.paragraph_format.space_before = Pt(5)
    table = _style(document, "Report Table", 9.5)
    table.paragraph_format.space_after = Pt(3)
    table.paragraph_format.space_before = Pt(3)
    table.paragraph_format.line_spacing = 1.1
    for name in ("Header", "Footer"):
        style = _style(document, name, 8.5)
        style.font.color.rgb = RGBColor.from_string(_MUTED)
    header = section.header.paragraphs[0]
    header.text = project_name or "Report"
    _rule(header, width="4")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ")
    _field(footer, "PAGE")
    footer.add_run(" of ")
    _field(footer, "NUMPAGES")
    document.settings.element.append(_element("w:updateFields", val="true"))
    properties = document.core_properties
    properties.title, properties.subject = title[:255], project_name
    properties.author = project_name or "Simon"
    properties.last_modified_by = "Simon report renderer " + RENDERER_VERSION
    # No invented document dates: source provenance is stored by the caller.
    for name in ("created", "modified"):
        element = properties._element.find("{http://purl.org/dc/terms/}" + name)
        if element is not None:
            properties._element.remove(element)
    properties.revision = 1
    properties.comments = "Formatted from the saved source; no content was generated."
    title_paragraph = document.add_paragraph(title, style="Title")
    _rule(title_paragraph, _ACCENT, width="14")
    return document


def _format(run: Run, *, bold: bool, italic: bool, strike: bool, code: bool) -> None:
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if strike:
        run.font.strike = True
    if code:
        _font_family(run._r.get_or_add_rPr(), "Courier New")
        run.font.size = Pt(9)
        run._r.get_or_add_rPr().append(_element("w:shd", fill=_PAPER, val="clear"))


def _hyperlink(paragraph: Paragraph, address: str) -> Any | None:
    # Preserve unsupported destinations as readable text; never create file/active links.
    if urlsplit(address).scheme.casefold() not in {"https", "http", "mailto"}:
        return None
    element = _element("w:hyperlink", history="1")
    identifier = paragraph.part.relate_to(address, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    element.set(qn("r:id"), identifier)
    paragraph._p.append(element)
    return element


def _inline(paragraph: Paragraph, tokens: list[Token], *, bold: bool = False) -> None:
    strong, emphasis, strike = int(bold), 0, 0
    link: Any | None = None
    unsupported_link: str | None = None

    def write(text: str, *, code: bool = False, destination: Any = None) -> None:
        run = paragraph.add_run(text)
        _format(run, bold=strong > 0, italic=emphasis > 0, strike=strike > 0, code=code)
        target = destination if destination is not None else link
        if target is not None:
            run.font.color.rgb = RGBColor.from_string(_ACCENT)
            run.underline = True
            target.append(run._r)

    for token in tokens:
        if token.type == "strong_open":
            strong += 1
        elif token.type == "strong_close":
            strong -= 1
        elif token.type == "em_open":
            emphasis += 1
        elif token.type == "em_close":
            emphasis -= 1
        elif token.type == "s_open":
            strike += 1
        elif token.type == "s_close":
            strike -= 1
        elif token.type == "link_open":
            address = str(token.attrGet("href") or "")
            link = _hyperlink(paragraph, address)
            unsupported_link = address if link is None else None
        elif token.type == "link_close":
            link = None
            if unsupported_link:
                write(" (" + unsupported_link + ")")
            unsupported_link = None
        elif token.type == "code_inline":
            write(token.content, code=True)
        elif token.type in {"softbreak", "hardbreak"}:
            write(" " if token.type == "softbreak" else "\n")
        elif token.type == "image":
            # No implicit remote fetch or arbitrary local file access while formatting.
            write(token.content or "Image")
            address = str(token.attrGet("src") or "")
            if address:
                write(" (")
                write(address, destination=_hyperlink(paragraph, address))
                write(")")
        elif token.type == "text" and link is None:
            position = 0
            for match in _URL.finditer(token.content):
                write(token.content[position : match.start()])
                address = match.group().rstrip(".,;:!?")
                for opening, closing in (("(", ")"), ("[", "]"), ("{", "}")):
                    while address.endswith(closing) and address.count(closing) > address.count(
                        opening
                    ):
                        address = address[:-1]
                write(address, destination=_hyperlink(paragraph, address))
                write(match.group()[len(address) :])
                position = match.end()
            write(token.content[position:])
        elif token.content:
            write(token.content)


def _numbering(document: DocumentType, *, ordered: bool, start: int, depth: int) -> int:
    numbering = document.part.numbering_part.element
    abstract_id = (
        max(
            (
                int(item.get(qn("w:abstractNumId")))
                for item in numbering.findall(qn("w:abstractNum"))
            ),
            default=-1,
        )
        + 1
    )
    abstract = _element("w:abstractNum", abstractNumId=str(abstract_id))
    abstract.append(_element("w:multiLevelType", val="singleLevel"))
    level = _element("w:lvl", ilvl="0")
    for child in (
        _element("w:start", val=str(start)),
        _element("w:numFmt", val="decimal" if ordered else "bullet"),
        _element("w:suff", val="tab"),
        _element("w:lvlText", val="%1." if ordered else "•"),
        _element("w:lvlJc", val="left"),
    ):
        level.append(child)
    properties = _element("w:pPr")
    indent = str((depth + 1) * 432)
    tabs = _element("w:tabs")
    tabs.append(_element("w:tab", val="num", pos=indent))
    properties.append(tabs)
    properties.append(_element("w:ind", left=indent, hanging="280"))
    level.append(properties)
    abstract.append(level)
    numbering.append(abstract)
    number = numbering.add_num(abstract_id)
    return int(number.numId)


@dataclass
class _List:
    number_id: int
    marker: bool = False


def _table(document: DocumentType, tokens: list[Token]) -> None:
    rows: list[list[Token]] = []
    for token in tokens:
        if token.type == "tr_open":
            rows.append([])
        elif token.type == "inline":
            rows[-1].append(token)
    columns = max((len(row) for row in rows), default=0)
    if not rows or not columns or columns > 20 or len(rows) > 1000:
        raise ValueError("Report table exceeds supported layout bounds")
    # Prose-heavy comparison matrices become full-width labeled records. This
    # retains every header/value while avoiding five narrow columns of tiny text.
    if columns > 6 or (
        columns > 4 and any(sum(len(cell.content) for cell in row) > 350 for row in rows[1:])
    ):
        for row in rows[1:]:
            for column, header in enumerate(rows[0]):
                if column >= len(row):
                    continue
                paragraph = document.add_paragraph(style="Heading 3" if column == 0 else "Normal")
                _inline(paragraph, header.children or [], bold=True)
                paragraph.add_run(": ")
                _inline(paragraph, row[column].children or [])
            spacer = document.add_paragraph()
            spacer.paragraph_format.space_after = Pt(3)
            _rule(spacer, width="4")
        return
    weights = [
        min(
            32,
            max(
                9, max((len(row[col].content) ** 0.5 * 3 if col < len(row) else 0) for row in rows)
            ),
        )
        for col in range(columns)
    ]
    widths = [Inches(7.1 * weight / sum(weights)) for weight in weights]
    table = document.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    properties = table._tbl.tblPr
    borders = _element("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        borders.append(_element("w:" + edge, val="single", sz="4", color=_LINE))
    for edge in ("left", "right", "insideV"):
        borders.append(_element("w:" + edge, val="nil"))
    properties.append(borders)
    margins = _element("w:tblCellMar")
    for side in ("top", "bottom", "left", "right"):
        margins.append(_element("w:" + side, w="85", type="dxa"))
    properties.append(margins)
    for index, width in enumerate(widths):
        table.columns[index].width = width
    for row_index, cells in enumerate(rows):
        row = table.rows[row_index]
        row._tr.get_or_add_trPr().append(_element("w:cantSplit"))
        if row_index == 0:
            row._tr.get_or_add_trPr().append(_element("w:tblHeader", val="true"))
        for column, cell in enumerate(row.cells):
            cell.width = widths[column]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_index == 0 or row_index % 2 == 0:
                cell._tc.get_or_add_tcPr().append(
                    _element(
                        "w:shd",
                        fill="EAEAE6" if row_index == 0 else "F7F7F5",
                        val="clear",
                    )
                )
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Report Table"]
            if column < len(cells):
                _inline(paragraph, cells[column].children or [], bold=row_index == 0)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _stable_package(document: DocumentType) -> bytes:
    original, result = io.BytesIO(), io.BytesIO()
    document.save(original)
    with ZipFile(original) as source, ZipFile(result, "w", ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            entry = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            entry.create_system = 0
            target.writestr(entry, source.read(name), compress_type=ZIP_DEFLATED, compresslevel=9)
    return result.getvalue()


def render_report(markdown: str, title: str, project_name: str = "") -> bytes:
    """Preserve saved prose as an editable, styled DOCX with deterministic bytes.

    Supports CommonMark plus tables/strikethrough. Images retain their alt text and
    source links without fetching resources. HTML stays literal text, never executes.
    Unsupported oversized inputs fail explicitly instead of silently dropping content.
    """
    if not isinstance(markdown, str) or not markdown.strip() or len(markdown) > MAX_MARKDOWN_CHARS:
        raise ValueError("Choose a nonempty report of at most 1000000 characters")
    if not title.strip() or len(title) > 300 or len(project_name) > 200:
        raise ValueError("Report title or project name exceeds its bound")
    if any(
        ord(character) < 32 and character not in "\n\r\t"
        for character in markdown + title + project_name
    ):
        raise ValueError("Report contains unsupported control characters")
    tokens = (
        MarkdownIt("commonmark", {"html": False})
        .enable(["table", "strikethrough"])
        .parse(prepare_report(markdown))
    )
    if len(tokens) > 100_000:
        raise ValueError("Report has too many structural elements")
    index = 0
    if tokens and tokens[0].type == "heading_open" and tokens[0].tag == "h1":
        # The source heading is authoritative; the caller's filename-derived title
        # is only a fallback. Render its original inline content once as the title.
        source_title = "".join(
            token.content
            for token in tokens[1].children or []
            if token.type in {"text", "code_inline", "image"}
        )
        if source_title:
            title = source_title
            index = 3
    document = _document(title, project_name)
    if index:
        title_paragraph = document.paragraphs[0]
        title_paragraph.text = ""
        _inline(title_paragraph, tokens[1].children or [])
    lists: list[_List] = []
    quote_depth = 0
    paragraph: Paragraph | None = None
    while index < len(tokens):
        token = tokens[index]
        if token.type == "table_open":
            end = next(
                position
                for position in range(index + 1, len(tokens))
                if tokens[position].type == "table_close"
            )
            _table(document, tokens[index : end + 1])
            index = end + 1
            continue
        if token.type in {"ordered_list_open", "bullet_list_open"}:
            if len(lists) >= 9:
                raise ValueError("Report lists exceed nine nesting levels")
            lists.append(
                _List(
                    _numbering(
                        document,
                        ordered=token.type == "ordered_list_open",
                        start=int(token.attrGet("start") or 1),
                        depth=len(lists),
                    )
                )
            )
        elif token.type in {"ordered_list_close", "bullet_list_close"}:
            lists.pop()
        elif token.type == "list_item_open":
            lists[-1].marker = True
        elif token.type == "blockquote_open":
            quote_depth += 1
        elif token.type == "blockquote_close":
            quote_depth -= 1
        elif token.type in {"paragraph_open", "heading_open"}:
            if token.type == "heading_open":
                paragraph = document.add_paragraph(style="Heading " + token.tag[1:])
            else:
                paragraph = document.add_paragraph(
                    style="Report Quote" if quote_depth else "Normal"
                )
            if lists:
                paragraph.paragraph_format.space_after = Pt(4)
                if lists[-1].marker:
                    numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                    numbering.get_or_add_ilvl().val = 0
                    numbering.get_or_add_numId().val = lists[-1].number_id
                    lists[-1].marker = False
                else:
                    paragraph.paragraph_format.left_indent = Inches(len(lists) * 432 / 1440)
            if quote_depth:
                borders = _element("w:pBdr")
                borders.append(_element("w:left", val="single", sz="14", color=_LINE, space="10"))
                paragraph._p.get_or_add_pPr().append(borders)
        elif token.type == "inline":
            if paragraph is None:
                raise ValueError("Report contains an unsupported block structure")
            _inline(paragraph, token.children or [])
        elif token.type in {"fence", "code_block"}:
            paragraph = document.add_paragraph(style="Report Code")
            paragraph._p.get_or_add_pPr().append(_element("w:shd", fill=_PAPER, val="clear"))
            paragraph.add_run(token.content)
        elif token.type == "hr":
            paragraph = document.add_paragraph()
            _rule(paragraph)
            paragraph.paragraph_format.space_after = Pt(8)
        elif token.content:
            document.add_paragraph(token.content)
        index += 1
    return _stable_package(document)
