import hashlib
import io
from uuid import NAMESPACE_URL, uuid4, uuid5

import pytest
from docx import Document

from simon.adapters.google import DRIVE_WRITE_SCOPE
from simon.adapters.project_drive import DriveError
from simon.domain.models import JobStatus, utc_now
from simon.domain.project_files import ProjectFileCreate
from simon.domain.project_outputs import PromoteProjectOutput
from simon.services.project_output_replication import ProjectOutputReplicationService
from simon.services.report_documents import DOCUMENT_KIND
from tests.contract.test_google_read_permissions import grant
from tests.contract.test_project_files import FakeDrive
from tests.contract.test_project_outputs import create_output as create_planning_output
from tests.contract.test_project_outputs import output_setup, replace_run


def create_output(h, number=1, *, count=1):
    run = create_planning_output(h, number, count=count)
    plan_job = h.store.get_job(run.plan_id)
    plan = h.platform.get(h.actor, run.plan_id)
    plan = plan.model_copy(
        update={"tasks": tuple(task.model_copy(update={"id": "draft"}) for task in plan.tasks)}
    )
    h.store.save_job(
        plan_job.model_copy(
            update={"input": {**plan_job.input, "plan": plan.model_dump(mode="json")}}
        ),
        plan_job.version,
    )
    saved = run.tasks[0].artifacts[0]
    answer = h.service.artifacts.publish_text(
        workspace_id=saved.workspace_id,
        actor_id=saved.actor_id,
        run_id=saved.run_id,
        task_id=saved.task_id,
        name="answer.txt",
        text="Result",
    )
    deliverable = h.service.artifacts.publish_bytes(
        workspace_id=answer.workspace_id,
        actor_id=answer.actor_id,
        run_id=answer.run_id,
        task_id=answer.task_id,
        name="Supplier comparison.md",
        media_type="text/markdown",
        content=b"Result",
    )
    run = run.model_copy(
        update={
            "tasks": tuple(
                task.model_copy(
                    update={
                        "id": "draft",
                        "output": "Result",
                        "artifacts": (answer, deliverable, *task.artifacts[1:]),
                    }
                )
                for task in run.tasks
            )
        }
    )
    replace_run(h, run)
    return run


@pytest.fixture
def replication(store, tmp_path):
    h = output_setup(store, tmp_path)
    grant(h.connected, h.actor, (DRIVE_WRITE_SCOPE,))
    h.cloud = h.connected.projects
    h.cloud.api = FakeDrive()
    h.replication = ProjectOutputReplicationService(h.service, h.cloud)
    h.cloud.output_sync = h.replication.sync
    return h


def test_team_outputs_replicate_after_local_commit_without_save_tool(replication):
    h = replication
    project_id = h.projects[0].id
    run = create_output(h, count=2)
    state = h.cloud.sync(h.actor, project_id, force=True)
    assert state["status"] == "ready"
    documents = {
        item["name"]: h.cloud.api.contents[identifier]
        for identifier, item in h.cloud.api.items.items()
        if item["name"].endswith(".docx")
    }
    assert len(documents) == 2
    assert "Result" in [
        p.text for p in Document(io.BytesIO(documents["Supplier comparison.docx"])).paragraphs
    ]
    assert "Deliverable 1/1" in [
        p.text for p in Document(io.BytesIO(documents["Report 1.docx"])).paragraphs
    ]
    assert h.replication.state(h.actor, project_id)["copied_outputs"] == 2
    names = {item["name"] for item in h.cloud.api.items.values()}
    assert {"Supplier comparison.docx", "Report 1.docx"} <= names
    assert "answer.txt" not in names
    assert h.service.read(h.actor, project_id, run.id, run.tasks[0].artifacts[0].id)[1] == b"Result"
    creates = h.cloud.api.creates
    for file_id in h.cloud.api.contents:
        if h.cloud.api.items[file_id]["name"] == "Supplier comparison.docx":
            h.cloud.api.contents[file_id] = b"User changed the remote copy"
    h.cloud.sync(h.actor, project_id, force=True)
    assert h.cloud.api.creates == creates
    assert b"User changed the remote copy" in h.cloud.api.contents.values()


def test_lost_upload_response_reuses_provider_id_after_reconstruction(replication):
    h = replication
    project_id = h.projects[0].id
    h.cloud.ensure(h.actor, project_id)
    run = create_output(h)
    h.cloud.api.timeout_after_create = True
    assert h.cloud.sync(h.actor, project_id, force=True)["status"] == "error"
    assert h.service.read(h.actor, project_id, run.id, run.tasks[0].artifacts[0].id)[1] == b"Result"
    assert h.cloud.api.creates == 2
    h.cloud.api.timeout_after_create = False
    restored = ProjectOutputReplicationService(h.service, h.cloud)
    h.cloud.output_sync = restored.sync
    assert h.cloud.sync(h.actor, project_id, force=True)["status"] == "ready"
    assert h.cloud.api.creates == 2
    assert restored.state(h.actor, project_id)["copied_outputs"] == 1


def test_drafts_stay_local_and_disabled_binding_prevents_replication(replication):
    h = replication
    project_id = h.projects[0].id
    run = create_output(h)
    replace_run(
        h,
        run.model_copy(
            update={
                "tasks": tuple(task.model_copy(update={"status": "failed"}) for task in run.tasks)
            }
        ),
    )
    assert h.cloud.sync(h.actor, project_id, force=True)["status"] == "ready"
    assert h.cloud.api.creates == 1
    create_output(h, 2)
    binding = h.cloud.binding(h.actor, project_id)
    h.store.save_project_drive(binding.model_copy(update={"enabled": False, "status": "unlinked"}))
    assert h.cloud.sync(h.actor, project_id, force=True)["status"] == "unlinked"
    assert h.cloud.api.creates == 1


def test_replication_cursor_survives_restart_and_batches_large_history(replication):
    h = replication
    project_id = h.projects[0].id
    create_output(h, count=7)
    assert h.cloud.sync(h.actor, project_id, force=True)["status"] == "pending"
    assert h.cloud.api.creates == 6
    h.cloud.output_sync = ProjectOutputReplicationService(h.service, h.cloud).sync
    binding = h.cloud.binding(h.actor, project_id)
    h.store.save_project_drive(binding.model_copy(update={"lease_until": utc_now()}))
    assert h.cloud.sync(h.actor, project_id, force=True)["status"] == "ready"
    assert h.cloud.api.creates == 8


def test_reviewed_candidate_stays_local_without_an_explicit_named_file(replication):
    h = replication
    run = create_output(h)
    executor = uuid4()
    h.runs.update(
        run.id,
        lambda current: current.model_copy(
            update={
                "executor_id": executor,
                "status": JobStatus.RUNNING,
                "tasks": (
                    current.tasks[0].model_copy(
                        update={"status": "running", "artifacts": (), "output": ""}
                    ),
                ),
            }
        ),
    )
    task = h.platform.get(h.actor, run.plan_id).tasks[0]
    arguments = {"executor_id": executor, "project_id": h.projects[0].id}
    text = "# Supplier report\nEvidence and recommendations."
    h.service.journal.append(
        h.actor, run.id, task, **arguments, kind="candidate", step=1, payload={"text": text}
    )
    h.service.journal.append(
        h.actor,
        run.id,
        task,
        **arguments,
        kind="review",
        step=2,
        payload={
            "status": "complete",
            "candidate_sha256": hashlib.sha256(text.encode()).hexdigest(),
        },
    )
    h.runs.update(
        run.id,
        lambda current: current.model_copy(
            update={
                "status": JobStatus.SUCCEEDED,
                "tasks": (
                    current.tasks[0].model_copy(update={"status": "succeeded", "output": text}),
                ),
            }
        ),
    )
    page = h.service.list(h.actor, h.projects[0].id)
    assert page.items[0].source == "candidate" and page.items[0].status == "accepted"
    assert page.items[0].project_copy is None
    assert h.cloud.sync(h.actor, h.projects[0].id, force=True)["status"] == "ready"
    assert text.encode() not in h.cloud.api.contents.values()
    assert h.cloud.api.creates == 1
    h.service.promote(
        h.actor,
        h.projects[0].id,
        run.id,
        page.items[0].id,
        PromoteProjectOutput(
            path="research/Supplier recommendations.md", idempotency_key="named-candidate-report"
        ),
        lambda: h.actor,
    )
    assert h.cloud.sync(h.actor, h.projects[0].id, force=True)["status"] == "ready"
    document_id = next(
        identifier
        for identifier, item in h.cloud.api.items.items()
        if item["name"] == "Supplier recommendations.docx"
    )
    report = Document(io.BytesIO(h.cloud.api.contents[document_id]))
    assert "Evidence and recommendations." in [paragraph.text for paragraph in report.paragraphs]


def test_named_response_publication_syncs_exact_report_without_random_prefix(replication):
    h = replication
    run = create_output(h)
    answer = run.tasks[0].artifacts[0]
    item = h.service.promote(
        h.actor,
        h.projects[0].id,
        run.id,
        answer.id,
        PromoteProjectOutput(
            path="research/Manufacturing economics.md", idempotency_key="publish-report"
        ),
        lambda: h.actor,
    )
    assert item.kind == "response"
    assert (
        h.replication.deliverable_name(h.actor, h.projects[0].id, item)
        == "Manufacturing economics.md"
    )
    assert h.cloud.sync(h.actor, h.projects[0].id, force=True)["status"] == "ready"
    assert {item["name"] for item in h.cloud.api.items.values()} >= {
        "Manufacturing economics.docx",
        "Supplier comparison.docx",
    }


def test_explicit_original_answer_copy_remains_discoverable_when_candidate_has_same_bytes(
    replication,
):
    h = replication
    run = create_output(h)
    answer = run.tasks[0].artifacts[0]
    executor = uuid4()
    h.runs.update(
        run.id,
        lambda current: current.model_copy(
            update={
                "status": JobStatus.RUNNING,
                "executor_id": executor,
            }
        ),
    )
    task = h.platform.get(h.actor, run.plan_id).tasks[0]
    arguments = {"executor_id": executor, "project_id": h.projects[0].id}
    h.service.journal.append(
        h.actor, run.id, task, **arguments, kind="candidate", step=1, payload={"text": "Result"}
    )
    h.service.journal.append(
        h.actor,
        run.id,
        task,
        **arguments,
        kind="review",
        step=2,
        payload={"status": "complete", "candidate_sha256": answer.sha256},
    )
    h.runs.update(
        run.id, lambda current: current.model_copy(update={"status": JobStatus.SUCCEEDED})
    )
    h.service.promote(
        h.actor,
        h.projects[0].id,
        run.id,
        answer.id,
        PromoteProjectOutput(
            path="research/Named complete report.md", idempotency_key="original-answer-report"
        ),
        lambda: h.actor,
    )
    assert h.cloud.sync(h.actor, h.projects[0].id, force=True)["status"] == "ready"
    assert "Named complete report.docx" in {item["name"] for item in h.cloud.api.items.values()}


@pytest.mark.parametrize("copy_state", ["generic", "edited", "missing"])
def test_generic_or_changed_response_copy_does_not_upload_original_bytes(replication, copy_state):
    h = replication
    run = create_output(h)
    answer = run.tasks[0].artifacts[0]
    request = PromoteProjectOutput(
        path=None if copy_state == "generic" else "reports/Research.md",
        idempotency_key="publish-report",
    )
    item = h.service.promote(h.actor, h.projects[0].id, run.id, answer.id, request, lambda: h.actor)
    copied = item.project_copy
    path = h.files.path(h.actor, copied.root, copied.path)
    if copy_state == "edited":
        path.write_text("User revisions after the original copy")
    elif copy_state == "missing":
        path.unlink()
    assert h.replication.deliverable_name(h.actor, h.projects[0].id, item) is None
    assert h.cloud.sync(h.actor, h.projects[0].id, force=True)["status"] == "ready"
    assert h.cloud.api.creates == 2  # Folder and the independently authored artifact.


def test_pending_old_prefixed_deliverable_reuses_original_receipt_and_provider_id(replication):
    h = replication
    project_id = h.projects[0].id
    binding = h.cloud.ensure(h.actor, project_id)
    run = create_output(h)
    artifact = run.tasks[0].artifacts[1]
    key = f"agent-output:{artifact.id}:{binding.folder_id}:{binding.google_email}"
    old_name = f"{str(artifact.id)[:8]}-{artifact.name}"[:200]
    h.cloud.api.timeout_after_create = True
    old = h.cloud.create_file(
        h.actor,
        ProjectFileCreate(project_id=project_id, name=old_name),
        key,
        lambda: h.actor,
        raw=b"Result",
        media_type=artifact.media_type,
        retry_create=True,
    )
    assert old["status"] == "unknown" and h.cloud.api.creates == 2
    h.cloud.api.timeout_after_create = False
    h.cloud.output_sync = ProjectOutputReplicationService(h.service, h.cloud).sync
    assert h.cloud.sync(h.actor, project_id, force=True)["status"] == "ready"
    receipt_id = uuid5(
        NAMESPACE_URL, f"project-file:{h.actor.workspace_id}:{h.actor.actor_id}:{key}"
    )
    receipt = h.store.project_file_operation(receipt_id)
    assert receipt.status == "succeeded" and receipt.file_id == old["file_id"]
    assert h.cloud.api.creates == 2
    assert (
        h.cloud.api.items[old["file_id"]]["name"] == old_name
    )  # Renaming is separate reviewed cleanup.
    h.cloud.sync(h.actor, project_id, force=True)
    assert h.cloud.api.creates == 2


def test_planning_and_controller_answers_never_sync_but_later_authored_answer_file_does(
    replication,
):
    h = replication
    create_planning_output(h, 3, count=2)
    run = create_output(h)
    source = run.tasks[0].artifacts[1]
    authored = h.service.artifacts.publish_bytes(
        workspace_id=source.workspace_id,
        actor_id=source.actor_id,
        run_id=source.run_id,
        task_id=source.task_id,
        name="answer.txt",
        media_type="text/plain",
        content=b"An actual exported file",
    )
    replace_run(
        h,
        run.model_copy(
            update={
                "tasks": (
                    run.tasks[0].model_copy(
                        update={
                            "artifacts": (*run.tasks[0].artifacts, authored),
                        }
                    ),
                )
            }
        ),
    )
    assert h.cloud.sync(h.actor, h.projects[0].id, force=True)["status"] == "ready"
    contents = [
        paragraph.text
        for content in h.cloud.api.contents.values()
        if content
        for paragraph in Document(io.BytesIO(content)).paragraphs
    ]
    assert "An actual exported file" in contents and "Result" in contents
    assert len(h.cloud.api.contents) == 3


def test_partial_page_failure_reconciles_copy_count_without_duplicates(replication, monkeypatch):
    h = replication
    project_id = h.projects[0].id
    h.cloud.ensure(h.actor, project_id)
    create_output(h, count=2)
    original_create = h.cloud.api.create

    def flaky_create(*args):
        result = original_create(*args)
        if args[1] == "Report 1.docx":
            raise DriveError("Lost response", unknown=True)
        return result

    monkeypatch.setattr(h.cloud.api, "create", flaky_create)
    assert h.cloud.sync(h.actor, project_id, force=True)["status"] == "error"
    assert h.cloud.api.creates == 3
    monkeypatch.setattr(h.cloud.api, "create", original_create)
    assert h.cloud.sync(h.actor, project_id, force=True)["status"] == "ready"
    assert h.cloud.api.creates == 3
    assert h.replication.state(h.actor, project_id)["copied_outputs"] == 2


def test_document_is_saved_before_cloud_call_with_distinct_stable_receipt(replication, monkeypatch):
    h = replication
    project = h.projects[0].id
    binding = h.cloud.ensure(h.actor, project)
    run = create_output(h)
    artifact = run.tasks[0].artifacts[1]
    original_create = h.cloud.api.create
    published = []

    def inspect_create(*args):
        assert args[1] == "Supplier comparison.docx"
        jobs = h.store.jobs(h.actor.workspace_id, h.actor.actor_id, DOCUMENT_KIND, 0, 100)
        assert len(jobs) == 1
        assert jobs[0].input["source"]["source_artifact_id"] == str(artifact.id)
        document, content = h.service.documents.materialize(h.actor, project, run.id, artifact.id)
        assert args[3] == content
        published.append(document)
        return original_create(*args)

    monkeypatch.setattr(h.cloud.api, "create", inspect_create)
    assert h.cloud.sync(h.actor, project, force=True)["status"] == "ready"
    key = h.service.documents.cloud_key(published[0], binding)
    receipt_id = uuid5(
        NAMESPACE_URL, f"project-file:{h.actor.workspace_id}:{h.actor.actor_id}:{key}"
    )
    assert h.store.project_file_operation(receipt_id).status == "succeeded"
    old_key = f"agent-output:{artifact.id}:{binding.folder_id}:{binding.google_email}"
    old_id = uuid5(
        NAMESPACE_URL, f"project-file:{h.actor.workspace_id}:{h.actor.actor_id}:{old_key}"
    )
    assert h.store.project_file_operation(old_id) is None
    assert h.cloud.sync(h.actor, project, force=True)["status"] == "ready"
    assert len(published) == 1


def test_successful_old_markdown_copy_is_not_replaced_or_duplicated(replication):
    h = replication
    project = h.projects[0].id
    binding = h.cloud.ensure(h.actor, project)
    run = create_output(h)
    artifact = run.tasks[0].artifacts[1]
    key = f"agent-output:{artifact.id}:{binding.folder_id}:{binding.google_email}"
    old = h.cloud.create_file(
        h.actor,
        ProjectFileCreate(project_id=project, name=artifact.name),
        key,
        lambda: h.actor,
        raw=b"Result",
        media_type=artifact.media_type,
        retry_create=True,
    )
    assert old["status"] == "succeeded"
    h.cloud.api.contents[old["file_id"]] = b"User edits retained"
    assert h.cloud.sync(h.actor, project, force=True)["status"] == "ready"
    assert h.cloud.api.creates == 2
    assert h.cloud.api.items[old["file_id"]]["name"] == "Supplier comparison.md"
    assert h.cloud.api.contents[old["file_id"]] == b"User edits retained"
