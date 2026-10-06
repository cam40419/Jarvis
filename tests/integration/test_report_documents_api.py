import io
from urllib.parse import unquote
from uuid import UUID, uuid4

from docx import Document

from simon.domain.agent_platform import AgentProfile, PlatformManifest, TeamTemplate
from simon.services.agent_platform import AgentPlatformService
from simon.services.document_rendering import DOCX_MEDIA_TYPE
from tests.contract.test_project_history import saved_run


def test_word_download_is_authorized_cached_and_preserves_immutable_source(
    client, container, auth_headers, tmp_path
):
    team = TeamTemplate(id="studio", name="Studio", agent_ids=("writer",))
    configured = AgentPlatformService(
        container.store,
        PlatformManifest(agents=(AgentProfile(id="writer", instructions="Write."),), teams=(team,)),
        state_dir=tmp_path / "agents",
        environ={},
    )
    container.agent_platform.__dict__.update(configured.__dict__)
    container.agent_platform.project_visibility_resolver = container.connected.projects.project
    container.agent_platform.project_team_resolver = lambda *_: team
    container.connected.settings = container.connected.settings.model_copy(
        update={"local_files_enabled": True, "local_files_dir": tmp_path / "local"}
    )
    project = UUID(
        client.post(
            "/v1/projects",
            headers=auth_headers,
            json={
                "name": "Formatted reports",
                "description": "Readable documents",
                "idempotency_key": "word-documents-api",
            },
        ).json()["id"]
    )
    run = saved_run(
        container.store, project, state_dir=container.agent_platform.state_dir, phase="execution"
    )
    artifact = run.tasks[0].artifacts[0]
    prefix = f"/v1/projects/{project}/outputs"
    listing = client.get(prefix).json()
    item = listing["items"][0]
    assert item["document_name"] == "Report.docx"
    assert item["document_url"] == f"{prefix}/{run.id}/{artifact.id}/document"
    response = client.get(item["document_url"] + "?title=untrusted-other-title")
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_MEDIA_TYPE
    assert (
        unquote(response.headers["content-disposition"])
        == "attachment; filename*=UTF-8''Report.docx"
    )
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    document = Document(io.BytesIO(response.content))
    assert "Result" in [paragraph.text for paragraph in document.paragraphs]
    assert "untrusted-other-title" not in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    assert client.get(item["document_url"]).content == response.content
    assert client.get(item["download_url"]).content == b"Result"
    assert (
        client.get(f"/v1/projects/{uuid4()}/outputs/{run.id}/{artifact.id}/document").status_code
        == 404
    )
    assert client.get(f"{prefix}/{run.id}/{uuid4()}/document").status_code == 404
    anonymous = client.__class__(client.app, base_url="http://localhost:8000")
    with anonymous:
        assert anonymous.get(item["document_url"]).status_code == 401
