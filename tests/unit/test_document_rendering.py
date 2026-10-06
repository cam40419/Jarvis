"""Saved reports retain their content and become deterministic, editable documents."""

import io
import socket
import zipfile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn

from simon.services.document_rendering import MAX_MARKDOWN_CHARS, render_report

REPORT = """# Manufacturing **Report**

## Executive summary

Use **supplier A** for *the pilot*. Unit economics: €18.90 \u00d7 120 = €2,268; margin 42.5%.
Read [the supplier source](https://example.com/prices?currency=USD&quantity=120).
The calculation uses `revenue - costs`; ~~superseded figure~~ remains visible.

3. Confirm the quote.
4. Inspect the sample.
   - Verify seams.
   - Check the label.

> Freight is not included. Request a written quote.

| Product | Cost | Decision |
| --- | ---: | --- |
| Shirt | $12.50 | **Pilot** |
| Hoodie | $31.20 | Defer |

```python
profit = revenue - costs
print("£27.50")
```

Direct source: https://example.com/products/(cotton).
"""


def all_text(document):
    return "\n".join(document.element.body.xpath(".//w:t/text()"))


def links(document):
    return {
        relation.target_ref
        for relation in document.part.rels.values()
        if relation.reltype == RELATIONSHIP_TYPE.HYPERLINK
    }


def test_semantic_structure_preserves_figures_prose_styles_and_clickable_sources():
    document = Document(io.BytesIO(render_report(REPORT, "Filename fallback", "Stdout Collective")))
    text = all_text(document)
    assert document.paragraphs[0].text == "Manufacturing Report"
    assert document.paragraphs[0].style.name == "Title"
    assert "Filename fallback" not in text
    assert [p.text for p in document.paragraphs if p.style.name == "Heading 2"] == [
        "Executive summary"
    ]
    assert "€18.90 \u00d7 120 = €2,268; margin 42.5%." in text
    assert {
        "https://example.com/prices?currency=USD&quantity=120",
        "https://example.com/products/(cotton)",
    } <= links(document)
    assert any(run.bold and "supplier A" in run.text for p in document.paragraphs for run in p.runs)
    assert any(
        run.italic and "the pilot" in run.text for p in document.paragraphs for run in p.runs
    )
    assert any(
        run.font.strike and "superseded figure" in run.text
        for p in document.paragraphs
        for run in p.runs
    )
    code = next(p for p in document.paragraphs if p.style.name == "Report Code")
    assert code.text == 'profit = revenue - costs\nprint("£27.50")\n'
    assert any(
        p.style.name == "Report Quote" and "Freight is not included" in p.text
        for p in document.paragraphs
    )
    table = document.tables[0]
    assert [[cell.text for cell in row.cells] for row in table.rows] == [
        ["Product", "Cost", "Decision"],
        ["Shirt", "$12.50", "Pilot"],
        ["Hoodie", "$31.20", "Defer"],
    ]
    assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert table.autofit is False
    assert any(run.bold and run.text == "Pilot" for run in table.cell(1, 2).paragraphs[0].runs)
    assert document.sections[0].header.paragraphs[0].text == "Stdout Collective"
    footer = document.sections[0].footer._element
    assert [value.strip() for value in footer.xpath(".//w:instrText/text()")] == [
        "PAGE",
        "NUMPAGES",
    ]


def test_ordered_and_nested_lists_are_editable_numbering_with_original_start():
    document = Document(io.BytesIO(render_report(REPORT, "Report")))
    paragraphs = [p for p in document.paragraphs if p._p.xpath("./w:pPr/w:numPr")]
    assert [p.text for p in paragraphs] == [
        "Confirm the quote.",
        "Inspect the sample.",
        "Verify seams.",
        "Check the label.",
    ]
    number_ids = [p._p.pPr.numPr.numId.val for p in paragraphs]
    assert number_ids[0] == number_ids[1] and number_ids[2] == number_ids[3]
    assert number_ids[0] != number_ids[2]
    numbering = document.part.numbering_part.element
    ordered = numbering.num_having_numId(number_ids[0]).abstractNumId.val
    element = next(
        e
        for e in numbering.findall(qn("w:abstractNum"))
        if int(e.get(qn("w:abstractNumId"))) == ordered
    )
    assert element.find(qn("w:lvl")).find(qn("w:start")).get(qn("w:val")) == "3"
    assert element.find(qn("w:lvl")).find(qn("w:numFmt")).get(qn("w:val")) == "decimal"
    assert element.find(qn("w:lvl")).find(qn("w:suff")).get(qn("w:val")) == "tab"
    properties = element.find(qn("w:lvl")).find(qn("w:pPr"))
    tab = properties.find(qn("w:tabs")).find(qn("w:tab"))
    assert tab.get(qn("w:val")) == "num"
    assert tab.get(qn("w:pos")) == properties.find(qn("w:ind")).get(qn("w:left"))


def test_wide_prose_table_becomes_readable_labeled_records_without_losing_cells():
    headers = ["Supplier", "Production", "Economics", "Risks", "Source"]
    cells = [
        "Supplier Alpha",
        "Production detail " * 18,
        "Cost is $18.90; MOQ is 120. " * 10,
        "Capacity requires confirmation. " * 9,
        "[Supplier page](https://example.com/supplier)",
    ]
    markdown = "| " + " | ".join(headers) + " |\n|" + "---|" * 5 + "\n| " + " | ".join(cells) + " |"
    document = Document(io.BytesIO(render_report(markdown, "Comparison")))
    assert not document.tables
    for header, value in zip(headers[:4], cells[:4], strict=True):
        assert any(p.text == header + ": " + value.strip() for p in document.paragraphs)
    assert "https://example.com/supplier" in links(document)
    assert "Supplier page" in all_text(document)


def test_compact_five_column_numeric_table_stays_a_table():
    markdown = (
        "| SKU | Cost | Price | Margin | Units |\n|---|---|---|---|---|\n"
        "| A | $12 | $30 | 60% | 120 |"
    )
    document = Document(io.BytesIO(render_report(markdown, "Unit economics")))
    assert len(document.tables) == 1 and len(document.tables[0].columns) == 5


def test_plain_report_outline_uses_source_title_and_recovers_headings():
    markdown = (
        "Competitive comparison report for Stdout Collective\n\nExecutive summary\n\n"
        "Market evidence is provisional.\n\n1) Supplier Alpha\n\nPositioning\n\n"
        "Positioning is documented.\n"
    )
    document = Document(io.BytesIO(render_report(markdown, "File name", "Stdout Collective")))
    assert document.paragraphs[0].text == "Competitive comparison report for Stdout Collective"
    assert sum(p.text == document.paragraphs[0].text for p in document.paragraphs) == 1
    assert any(
        p.text == "1) Supplier Alpha" and p.style.name == "Heading 2" for p in document.paragraphs
    )
    assert "Market evidence is provisional." in all_text(document)


def test_same_source_is_byte_identical_even_when_packaging_clock_changes(monkeypatch):
    monkeypatch.setattr(zipfile.time, "localtime", lambda *_: (2025, 1, 1, 10, 0, 0, 0, 1, 0))
    first = render_report(REPORT, "Report", "Stdout Collective")
    monkeypatch.setattr(zipfile.time, "localtime", lambda *_: (2030, 12, 25, 22, 30, 0, 0, 1, 0))
    second = render_report(REPORT, "Report", "Stdout Collective")
    assert first == second
    with zipfile.ZipFile(io.BytesIO(first)) as package:
        assert all(item.date_time == (1980, 1, 1, 0, 0, 0) for item in package.infolist())
        assert b"dcterms:created" not in package.read("docProps/core.xml")
        assert b"dcterms:modified" not in package.read("docProps/core.xml")


@pytest.mark.parametrize(
    "source",
    [REPORT, "Competitive comparison report\n\nScope and method\nRead the supplied sources."],
)
def test_explicit_fonts_and_single_title_rule_survive_source_title_replacement(source):
    document = Document(io.BytesIO(render_report(source, "Report")))
    for name in (
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Report Table",
        "Header",
        "Footer",
        "Report Code",
    ):
        fonts = document.styles[name].element.rPr.rFonts
        assert not any(attribute.casefold().endswith("theme") for attribute in fonts.attrib)
        family = "Courier New" if name == "Report Code" else "Arial"
        assert all(
            fonts.get(qn("w:" + slot)) == family for slot in ("ascii", "hAnsi", "eastAsia", "cs")
        )
    borders = document.paragraphs[0]._p.xpath("./w:pPr/w:pBdr/w:bottom")
    assert len(borders) == 1 and borders[0].get(qn("w:color")) == "315D8A"
    assert not document.styles["Title"].element.xpath("./w:pPr/w:pBdr")


def test_images_and_html_cannot_fetch_resources_or_create_active_file_links(monkeypatch):
    monkeypatch.setattr(
        socket, "create_connection", lambda *_: pytest.fail("No network during rendering")
    )
    source = '![Pattern](https://example.com/pattern.png)\n\n[Local file](file:///private/key)\n\n<script>alert("literal")</script>'
    result = render_report(source, "Safe report")
    document = Document(io.BytesIO(result))
    assert "Pattern" in all_text(document)
    assert '<script>alert("literal")</script>' in all_text(document)
    assert all(url.startswith(("https://", "http://", "mailto:")) for url in links(document))
    with zipfile.ZipFile(io.BytesIO(result)) as package:
        assert not any(name.startswith("word/media/") for name in package.namelist())


@pytest.mark.parametrize(
    "source",
    ["", " ", "x" * (MAX_MARKDOWN_CHARS + 1), "bad\x00control"],
    ids=["empty", "blank", "oversized", "control"],
)
def test_invalid_or_oversized_reports_fail_instead_of_truncating(source):
    with pytest.raises(ValueError):
        render_report(source, "Report")
