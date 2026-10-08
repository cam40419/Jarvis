"""Source originals stay immutable while extraction remains bounded and inert."""

import hashlib
import io
import os
from concurrent.futures import ThreadPoolExecutor
from uuid import uuid4

import pytest
from docx import Document

from simon.domain.errors import ValidationError
from simon.services.intake_sources import (
    MAX_SOURCE_BYTES,
    TEXT_LIMIT,
    IntakeSourceBytes,
    extract_source,
    redact,
)


@pytest.mark.parametrize(
    ("name", "mime", "content"),
    [
        ("notes.txt", "text/plain", b"Project notes."),
        ("brief.md", "application/octet-stream", b"# Project\nEvidence."),
        ("budget.csv", "application/octet-stream", b"line,cost\nA,10"),
        ("metadata.json", "application/json", b'{"source":"owner"}'),
        ("memo", "text/plain", b"\xef\xbb\xbfUnicode BOM."),
        ("unsafe.html", "text/html", b'<script>fetch("https://evil.example")</script>'),
        ("notes.xml", "application/xml", b'<!DOCTYPE foo [<!ENTITY external SYSTEM "file:///x">]>'),
    ],
)
def test_text_extraction_is_inert_and_preserves_original_bytes(name, mime, content):
    original = bytes(content)
    result = extract_source(content, name, mime)
    assert result.status == "text"
    assert result.text == content.decode("utf-8-sig")
    assert content == original
    assert not result.truncated and not result.redactions


@pytest.mark.parametrize(
    ("name", "mime", "content"),
    [
        ("photo.png", "image/png", b"\x89PNG\r\n"),
        ("scan.pdf", "application/pdf", b"%PDF-1.4 fake"),
        ("table.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK"),
        ("binary.txt", "text/plain", b"\xff\xfe\x01"),
        ("nul.txt", "text/plain", b"Safe\x00not text"),
        ("broken.docx", "application/octet-stream", b"not a DOCX"),
        ("archive.zip", "application/zip", b"PK\x03\x04"),
    ],
)
def test_unknown_or_malformed_files_are_explicitly_unparsed(name, mime, content):
    result = extract_source(content, name, mime)
    assert result.status == "unparsed" and result.text == ""
    assert not result.truncated and not result.redactions


def test_docx_extraction_uses_bounded_package_reader_for_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Evidence-backed brand direction.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Essential"
    table.cell(0, 1).text = "T-shirt"
    output = io.BytesIO()
    document.save(output)
    result = extract_source(output.getvalue(), "brand.docx", "application/octet-stream")
    assert result.status == "text"
    assert "Evidence-backed brand direction." in result.text
    assert "Essential" in result.text and "T-shirt" in result.text


@pytest.mark.parametrize(
    "secret",
    [
        "sk-" + "a" * 30,
        "AIza" + "b" * 32,
        "ghp_" + "c" * 25,
        "Bearer " + "d" * 25,
        "API_KEY=arbitrary-secret-value",
        'password: "arbitrary-secret-value"',
        "-----BEGIN PRIVATE KEY-----\nprivate material\n-----END PRIVATE KEY-----",
    ],
)
def test_credential_redaction_happens_before_model_excerpting(secret):
    result = extract_source(f"Before\n{secret}\nAfter".encode(), "notes.txt", "text/plain")
    assert secret not in result.text
    assert "[REDACTED CREDENTIAL]" in result.text
    assert result.redactions >= 1 and result.status == "text"
    assert "Before" in result.text and "After" in result.text
    assert redact(result.text) == (result.text, 0)


def test_extraction_bounds_count_redactions_beyond_visible_excerpt():
    content = "x" * (TEXT_LIMIT + 100) + "\nAPI_KEY=private-beyond-excerpt"
    result = extract_source(content.encode(), "long.txt", "text/plain")
    assert result.truncated and len(result.text) == TEXT_LIMIT
    assert result.redactions == 1


def test_blob_storage_is_immutable_deduplicated_and_scoped_by_project(tmp_path):
    blobs = IntakeSourceBytes(tmp_path / "blobs")
    workspace, project, other_project = uuid4(), uuid4(), uuid4()
    content = b"Original archive evidence, unchanged."
    sha = blobs.put(workspace, project, content)
    assert sha == hashlib.sha256(content).hexdigest()
    assert blobs.put(workspace, project, content) == sha
    assert blobs.read(workspace, project, sha, len(content)) == content
    assert len(list((tmp_path / "blobs" / str(workspace) / str(project)).iterdir())) == 1
    with pytest.raises(ValidationError, match="unavailable"):
        blobs.read(workspace, other_project, sha, len(content))
    assert blobs.put(workspace, other_project, content) == sha
    assert len(list((tmp_path / "blobs" / str(workspace)).iterdir())) == 2


def test_concurrent_duplicate_publication_keeps_one_integrity_checked_file(tmp_path):
    blobs = IntakeSourceBytes(tmp_path / "blobs")
    workspace, project = uuid4(), uuid4()
    content = b"One immutable source." * 100
    with ThreadPoolExecutor(max_workers=8) as pool:
        hashes = list(pool.map(lambda _: blobs.put(workspace, project, content), range(24)))
    assert len(set(hashes)) == 1
    assert blobs.read(workspace, project, hashes[0], len(content)) == content
    assert not list((tmp_path / "blobs").rglob(".upload-*"))


def test_blob_read_detects_content_and_size_corruption_and_refuses_overwrite(tmp_path):
    blobs = IntakeSourceBytes(tmp_path / "blobs")
    workspace, project = uuid4(), uuid4()
    raw = b"original"
    sha = blobs.put(workspace, project, raw)
    with pytest.raises(ValidationError, match="integrity"):
        blobs.read(workspace, project, sha, len(raw) + 1)
    target = tmp_path / "blobs" / str(workspace) / str(project) / sha
    target.write_bytes(b"tampered")
    with pytest.raises(ValidationError, match="integrity"):
        blobs.read(workspace, project, sha, len(raw))
    with pytest.raises(ValidationError, match="integrity"):
        blobs.put(workspace, project, raw)
    assert target.read_bytes() == b"tampered"


@pytest.mark.parametrize("digest", ["../secret", "A" * 64, "0" * 63, "0" * 65, ""])
def test_invalid_blob_identifier_cannot_become_a_path(tmp_path, digest):
    blobs = IntakeSourceBytes(tmp_path / "blobs")
    with pytest.raises(ValidationError, match="digest"):
        blobs.read(uuid4(), uuid4(), digest, 0)
    assert not (tmp_path / "blobs").exists()


def test_blob_size_limit_precedes_storage_mutations(tmp_path):
    blobs = IntakeSourceBytes(tmp_path / "blobs")
    with pytest.raises(ValidationError, match="limit"):
        blobs.put(uuid4(), uuid4(), b"x" * (MAX_SOURCE_BYTES + 1))
    assert not (tmp_path / "blobs").exists()


def test_publication_failure_cleans_owned_staging_file(tmp_path, monkeypatch):
    blobs = IntakeSourceBytes(tmp_path / "blobs")

    def refuse_link(*_args, **_kwargs):
        raise OSError("Synthetic filesystem publication failure")

    monkeypatch.setattr(os, "link", refuse_link)
    with pytest.raises((OSError, ValidationError)):
        blobs.put(uuid4(), uuid4(), b"evidence")
    assert not [path for path in (tmp_path / "blobs").rglob("*") if path.is_file()]


def test_filesystem_redirect_cannot_expose_another_storage_tree(tmp_path):
    actual, alias = tmp_path / "actual", tmp_path / "alias"
    actual.mkdir()
    try:
        alias.symlink_to(actual, target_is_directory=True)
    except OSError:
        pytest.skip("Creating symlinks is unavailable on this host")
    blobs = IntakeSourceBytes(alias)
    with pytest.raises(ValidationError, match="redirect"):
        blobs.put(uuid4(), uuid4(), b"evidence")
    assert list(actual.iterdir()) == []


def test_unterminated_private_key_redacts_remainder_instead_of_exposing_truncated_material():
    content = "Public context.\n-----BEGIN RSA PRIVATE KEY-----\nprivate-incomplete-material"
    result = extract_source(content.encode(), "source.txt", "text/plain")
    assert result.text == "Public context.\n[REDACTED CREDENTIAL]"
    assert result.redactions == 1 and not result.truncated
    assert "private-incomplete-material" not in result.text


def test_repeated_unterminated_key_markers_collapse_to_one_redacted_region():
    content = "Public context.\n" + (
        "-----BEGIN PRIVATE KEY-----\nprivate-repeated-material\n" * 2048
    )
    result = extract_source(content.encode(), "source.txt", "text/plain")
    assert result.text == "Public context.\n[REDACTED CREDENTIAL]"
    assert result.redactions == 1 and not result.truncated


def test_independent_complete_and_incomplete_keys_preserve_only_public_surrounding_text():
    content = (
        "First public.\n-----BEGIN PRIVATE KEY-----\nprivate-one\n-----END PRIVATE KEY-----"
        "\nSecond public.\n-----BEGIN EC PRIVATE KEY-----\nprivate-two"
    )
    cleaned, count = redact(content)
    assert cleaned == "First public.\n[REDACTED CREDENTIAL]\nSecond public.\n[REDACTED CREDENTIAL]"
    assert count == 2
